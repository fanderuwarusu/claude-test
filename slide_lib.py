"""
slide_lib.py  ─  AI講義スライド生成ライブラリ（固定レンダラー層）

【使い方】
    from slide_lib import build_pptx, build_pdf, SLIDE_TYPES

    SLIDE_DEFS = [
        {"type": "title",   "title": "...", "subtitle": "...", "label": "..."},
        {"type": "agenda",  "title": "...", "items": [("0", "説明", "5分"), ...]},
        {"type": "section", "num": 1, "title": "...", "subtitle": "..."},
        {"type": "content", "title": "...", "bullets": ["..."], "ph": False, "ph_label": "...", "note": ""},
        {"type": "two_col", "title": "...",
            "left_title": "...", "right_title": "...",
            "left_bullets": ["..."], "right_bullets": ["..."]},
        {"type": "work",    "title": "...", "instruction": "...", "work_label": "ワーク（X分）"},
    ]

    build_pptx(SLIDE_DEFS, "出力ファイル名")   # .pptx を生成
    build_pdf(SLIDE_DEFS,  "出力ファイル名")   # .pdf  を生成

【SLIDE_DEFS フィールド一覧】
    title   : title / subtitle / label
    agenda  : title / items=[("番号", "ラベル", "時間"), ...]
    section : num / title / subtitle(省略可)
    content : title / bullets=["##見出し" or "本文",...] / ph(bool) / ph_label / note
    two_col : title / left_title / right_title / left_bullets / right_bullets
              ※ left_bullets/right_bullets を省略するとプレースホルダーになる
    work    : title / instruction / work_label

    bullets で "##" から始まる行は見出し扱い（色・サイズが変わる）
"""

# ============================================================
# 共通カラー定数（テーマA ネイビー × ブルー）
# ============================================================
_THEME = {
    "primary": (0x1A, 0x2E, 0x4A),
    "accent":  (0x00, 0x8C, 0xD7),
    "bg":      (0xF0, 0xF6, 0xFC),
    "white":   (0xFF, 0xFF, 0xFF),
    "gray":    (0x6B, 0x7A, 0x8D),
    "ph_c":    (0xCC, 0xCC, 0xCC),
    "ph_bg":   (0xF7, 0xF7, 0xF7),
    "lt_blue": (0xA0, 0xC4, 0xE8),
    "row_alt": (0xD8, 0xEE, 0xFA),
}

# 利用可能なスライドタイプ一覧（Skill参照用）
SLIDE_TYPES = ["title", "agenda", "section", "content", "two_col", "work"]


# ============================================================
# ─── PPTX レンダラー ────────────────────────────────────────
# ============================================================
def build_pptx(slide_defs: list, filename: str) -> str:
    """SLIDE_DEFS からPPTXを生成し保存する。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    W = Inches(10.0)
    H = Inches(7.5)

    C_PRIMARY = RGBColor(*_THEME["primary"])
    C_ACCENT  = RGBColor(*_THEME["accent"])
    C_BG      = RGBColor(*_THEME["bg"])
    C_WHITE   = RGBColor(*_THEME["white"])
    C_GRAY    = RGBColor(*_THEME["gray"])
    C_PH      = RGBColor(*_THEME["ph_c"])
    C_LT_BLUE = RGBColor(*_THEME["lt_blue"])
    C_ROW_ALT = RGBColor(*_THEME["row_alt"])

    # ── ユーティリティ ─────────────────────────────────────
    def bg(slide, color):
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = color

    def rect(slide, x, y, w, h, color):
        s = slide.shapes.add_shape(1, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()
        return s

    def txt(slide, s, x, y, w, h, size=14, bold=False, color=None,
            align=PP_ALIGN.LEFT):
        if color is None:
            color = C_PRIMARY
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = s
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return tb

    def ph(slide, x, y, w, h, label="※ここに内容を入力"):
        r = slide.shapes.add_shape(1, x, y, w, h)
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor(*_THEME["ph_bg"])
        r.line.color.rgb = C_PH
        r.line.width = Pt(1)
        tf = r.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(12)
        run.font.color.rgb = C_PH
        run.font.italic = True
        return r

    def header(slide, title):
        rect(slide, Inches(0), Inches(0), W, Inches(1.0), C_PRIMARY)
        rect(slide, Inches(0), Inches(1.0), W, Inches(0.06), C_ACCENT)
        txt(slide, title, Inches(0.35), Inches(0.15), Inches(9.3), Inches(0.75),
            size=22, bold=True, color=C_WHITE)

    # ── スライド描画 ────────────────────────────────────────
    def make_title(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_PRIMARY)
        rect(sl, Inches(0), Inches(0), Inches(0.45), H, C_ACCENT)
        txt(sl, d.get("label", ""),
            Inches(0.7), Inches(1.4), Inches(8.8), Inches(0.55),
            size=14, bold=True, color=C_ACCENT)
        txt(sl, d["title"],
            Inches(0.7), Inches(2.1), Inches(9.0), Inches(1.5),
            size=36, bold=True, color=C_WHITE)
        txt(sl, d.get("subtitle", ""),
            Inches(0.7), Inches(3.75), Inches(9.0), Inches(0.8),
            size=18, color=C_LT_BLUE)
        txt(sl, "【日付】　【講師名】",
            Inches(0.7), Inches(5.6), Inches(7), Inches(0.45),
            size=13, color=C_GRAY)

    def make_agenda(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_BG)
        header(sl, d["title"])
        y = Inches(1.18)
        rh = Inches(0.72)
        for i, (num, label, dur) in enumerate(d["items"]):
            clr = C_ACCENT if i % 2 == 0 else C_PRIMARY
            rect(sl, Inches(0.35), y, Inches(0.6), rh, clr)
            txt(sl, num, Inches(0.35), y + Inches(0.15), Inches(0.6), Inches(0.42),
                size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
            rect(sl, Inches(0.97), y, Inches(7.15), rh, C_WHITE)
            txt(sl, label, Inches(1.07), y + Inches(0.15), Inches(7.0), Inches(0.42),
                size=13, color=C_PRIMARY)
            rect(sl, Inches(8.15), y, Inches(1.5), rh, C_ROW_ALT)
            txt(sl, dur, Inches(8.15), y + Inches(0.15), Inches(1.5), Inches(0.42),
                size=12, color=C_ACCENT, align=PP_ALIGN.CENTER)
            y += rh + Inches(0.04)

    def make_section(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_ACCENT)
        rect(sl, Inches(0), Inches(0), W, Inches(0.07), C_WHITE)
        rect(sl, Inches(0), Inches(7.43), W, Inches(0.07), C_WHITE)
        txt(sl, f"Section {d['num']}",
            Inches(0), Inches(2.3), W, Inches(0.65),
            size=18, color=C_WHITE, align=PP_ALIGN.CENTER)
        txt(sl, d["title"],
            Inches(0.5), Inches(2.95), Inches(9.0), Inches(1.1),
            size=32, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        if d.get("subtitle"):
            txt(sl, d["subtitle"],
                Inches(0.5), Inches(4.15), Inches(9.0), Inches(0.6),
                size=16, color=C_LT_BLUE, align=PP_ALIGN.CENTER)

    def make_content(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_BG)
        header(sl, d["title"])
        bullets = d.get("bullets")
        if bullets:
            tb = sl.shapes.add_textbox(Inches(0.4), Inches(1.2),
                                       Inches(9.2), Inches(5.5))
            tf = tb.text_frame
            tf.word_wrap = True
            first = True
            for b in bullets:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.space_before = Pt(4)
                p.alignment = PP_ALIGN.LEFT
                run = p.add_run()
                if b.startswith("##"):
                    run.text = b[2:].strip()
                    run.font.size = Pt(15)
                    run.font.bold = True
                    run.font.color.rgb = C_ACCENT
                else:
                    run.text = f"▶  {b}"
                    run.font.size = Pt(14)
                    run.font.color.rgb = C_PRIMARY
        if d.get("ph"):
            ph_y = Inches(5.4) if bullets else Inches(1.3)
            ph_h = Inches(1.6) if bullets else Inches(5.5)
            ph(sl, Inches(0.4), ph_y, Inches(9.2), ph_h,
               d.get("ph_label", "※ここに内容を入力"))
        if d.get("note"):
            txt(sl, f"📌 {d['note']}",
                Inches(0.4), Inches(7.05), Inches(9.2), Inches(0.35),
                size=10, color=C_GRAY)

    def make_two_col(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_BG)
        header(sl, d["title"])
        cw = Inches(4.4)
        lx = Inches(0.35)
        rx = lx + cw + Inches(0.3)
        rect(sl, lx, Inches(1.18), cw, Inches(0.42), C_ACCENT)
        txt(sl, d["left_title"],
            lx + Inches(0.08), Inches(1.2), cw - Inches(0.16), Inches(0.38),
            size=13, bold=True, color=C_WHITE)
        rect(sl, rx, Inches(1.18), cw, Inches(0.42), C_PRIMARY)
        txt(sl, d["right_title"],
            rx + Inches(0.08), Inches(1.2), cw - Inches(0.16), Inches(0.38),
            size=13, bold=True, color=C_WHITE)
        for x, key in [(lx, "left_bullets"), (rx, "right_bullets")]:
            items = d.get(key)
            if items:
                tb = sl.shapes.add_textbox(x + Inches(0.1), Inches(1.72),
                                           cw - Inches(0.2), Inches(5.3))
                tf = tb.text_frame
                tf.word_wrap = True
                first = True
                for b in items:
                    p2 = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p2.space_before = Pt(3)
                    run = p2.add_run()
                    run.text = f"・{b}"
                    run.font.size = Pt(13)
                    run.font.color.rgb = C_PRIMARY
            else:
                ph(sl, x, Inches(1.72), cw, Inches(5.3), "※内容を入力")

    def make_work(prs, d):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg(sl, C_PRIMARY)
        rect(sl, Inches(0.4), Inches(0.3), Inches(1.7), Inches(0.55), C_ACCENT)
        txt(sl, f"✏  {d.get('work_label', 'ワーク')}",
            Inches(0.45), Inches(0.32), Inches(1.6), Inches(0.5),
            size=13, bold=True, color=C_WHITE)
        txt(sl, d["title"],
            Inches(0.4), Inches(1.1), Inches(9.2), Inches(1.0),
            size=26, bold=True, color=C_WHITE)
        txt(sl, d.get("instruction", ""),
            Inches(0.4), Inches(2.2), Inches(9.2), Inches(0.65),
            size=15, color=C_LT_BLUE)
        ph(sl, Inches(0.4), Inches(3.05), Inches(9.2), Inches(3.9),
           "※ワーク内容・手順をここに記載")

    _RENDERERS = {
        "title":   make_title,
        "agenda":  make_agenda,
        "section": make_section,
        "content": make_content,
        "two_col": make_two_col,
        "work":    make_work,
    }

    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    for d in slide_defs:
        _RENDERERS[d["type"]](prs, d)

    out = f"{filename}.pptx" if not filename.endswith(".pptx") else filename
    prs.save(out)
    print(f"✅ PPTX生成完了: {out}  ({len(prs.slides)}枚)")
    return out


# ============================================================
# ─── PDF レンダラー ─────────────────────────────────────────
# ============================================================
def build_pdf(slide_defs: list, filename: str) -> str:
    """SLIDE_DEFS からPDFを生成し保存する。"""
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.colors import HexColor

    FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
    pdfmetrics.registerFont(TTFont("JA", FONT_PATH))
    F = "JA"

    PW = 10 * inch
    PH = 7.5 * inch

    def _hex(rgb):
        return HexColor("#{:02X}{:02X}{:02X}".format(*rgb))

    PRIMARY  = _hex(_THEME["primary"])
    ACCENT   = _hex(_THEME["accent"])
    BG       = _hex(_THEME["bg"])
    WHITE    = _hex(_THEME["white"])
    GRAY     = _hex(_THEME["gray"])
    PH_C     = _hex(_THEME["ph_c"])
    PH_BG    = _hex(_THEME["ph_bg"])
    LT_BLUE  = _hex(_THEME["lt_blue"])
    ROW_ALT  = _hex(_THEME["row_alt"])

    # ── ユーティリティ ─────────────────────────────────────
    def fill_bg(c, color=None):
        c.setFillColor(color or BG)
        c.rect(0, 0, PW, PH, fill=1, stroke=0)

    def drect(c, x, y, w, h, color):
        c.setFillColor(color)
        c.rect(x, PH - y - h, w, h, fill=1, stroke=0)

    def dtxt(c, s, x, y, size=14, color=PRIMARY, align="left", max_w=None):
        c.setFillColor(color)
        c.setFont(F, size)
        ry = PH - y - size
        if align == "center":
            c.drawCentredString(x + (max_w or 0) / 2, ry, s)
        elif align == "right":
            c.drawRightString(x + (max_w or 0), ry, s)
        else:
            c.drawString(x, ry, s)

    def dph(c, x, y, w, h, label="※ここに内容を入力"):
        drect(c, x, y, w, h, PH_BG)
        c.setStrokeColor(PH_C)
        c.setLineWidth(0.5)
        c.rect(x, PH - y - h, w, h, fill=0, stroke=1)
        dtxt(c, label, x, y + h / 2 - 6, size=11, color=PH_C,
             align="center", max_w=w)

    def header_bar(c, title):
        drect(c, 0, 0, PW, 76, PRIMARY)
        drect(c, 0, 76, PW, 4, ACCENT)
        dtxt(c, title, 25, 14, size=20, color=WHITE)

    # ── スライド描画 ────────────────────────────────────────
    def render_title(c, d):
        fill_bg(c, PRIMARY)
        drect(c, 0, 0, 32, PH, ACCENT)
        dtxt(c, d.get("label", ""), 52, 102, size=13, color=ACCENT)
        dtxt(c, d["title"],         52, 152, size=34, color=WHITE)
        dtxt(c, d.get("subtitle", ""), 52, 270, size=16, color=LT_BLUE)
        dtxt(c, "【日付】　【講師名】", 52, 408, size=12, color=GRAY)

    def render_agenda(c, d):
        fill_bg(c)
        header_bar(c, d["title"])
        y = 90
        for i, (num, label, dur) in enumerate(d["items"]):
            bg_ = ACCENT if i % 2 == 0 else PRIMARY
            drect(c, 25, y, 42, 46, bg_)
            dtxt(c, num, 25, y + 13, size=12, color=WHITE,
                 align="center", max_w=42)
            drect(c, 69, y, 523, 46, WHITE)
            dtxt(c, label, 77, y + 14, size=12, color=PRIMARY)
            drect(c, 594, y, 106, 46, ROW_ALT)
            dtxt(c, dur, 594, y + 14, size=11, color=ACCENT,
                 align="center", max_w=106)
            y += 52

    def render_section(c, d):
        fill_bg(c, ACCENT)
        drect(c, 0, 0, PW, 5, WHITE)
        drect(c, 0, PH - 5, PW, 5, WHITE)
        dtxt(c, f"Section {d['num']}", 0, 166, size=16, color=WHITE,
             align="center", max_w=PW)
        dtxt(c, d["title"], 38, 206, size=28, color=WHITE,
             align="center", max_w=PW - 76)
        if d.get("subtitle"):
            dtxt(c, d["subtitle"], 38, 300, size=14, color=LT_BLUE,
                 align="center", max_w=PW - 76)

    def render_content(c, d):
        fill_bg(c)
        header_bar(c, d["title"])
        y = 94
        bullets = d.get("bullets")
        if bullets:
            for b in bullets:
                if b.startswith("##"):
                    dtxt(c, b[2:].strip(), 28, y, size=14, color=ACCENT)
                    y += 22
                else:
                    dtxt(c, f"▶  {b}", 28, y, size=13, color=PRIMARY)
                    y += 20
                y += 4
        if d.get("ph"):
            ph_y = y + 6 if bullets else 94
            ph_h = min(120, PH - ph_y - 30) if bullets else (PH - 94 - 30)
            dph(c, 28, ph_y, PW - 56, ph_h, d.get("ph_label", "※ここに内容を入力"))
        if d.get("note"):
            dtxt(c, f"📌 {d['note']}", 28, PH - 30, size=9, color=GRAY)

    def render_two_col(c, d):
        fill_bg(c)
        header_bar(c, d["title"])
        cw = 308
        lx, rx = 25, 25 + cw + 22
        drect(c, lx, 85, cw, 30, ACCENT)
        dtxt(c, d["left_title"],  lx + 6, 90, size=12, color=WHITE)
        drect(c, rx, 85, cw, 30, PRIMARY)
        dtxt(c, d["right_title"], rx + 6, 90, size=12, color=WHITE)
        for x, key in [(lx, "left_bullets"), (rx, "right_bullets")]:
            items = d.get(key)
            if items:
                y = 124
                for b in items:
                    dtxt(c, f"・{b}" if b else "", x + 6, y, size=12, color=PRIMARY)
                    y += 18
            else:
                dph(c, x, 118, cw, PH - 138, "※内容を入力")

    def render_work(c, d):
        fill_bg(c, PRIMARY)
        drect(c, 28, 21, 120, 38, ACCENT)
        dtxt(c, f"✏  {d.get('work_label', 'ワーク')}", 34, 27, size=12, color=WHITE)
        dtxt(c, d["title"],             28, 80,  size=24, color=WHITE)
        dtxt(c, d.get("instruction", ""), 28, 158, size=13, color=LT_BLUE)
        dph(c, 28, 182, PW - 56, PH - 210, "※ワーク内容・手順をここに記載")

    _RENDERERS = {
        "title":   render_title,
        "agenda":  render_agenda,
        "section": render_section,
        "content": render_content,
        "two_col": render_two_col,
        "work":    render_work,
    }

    out = f"{filename}.pdf" if not filename.endswith(".pdf") else filename
    c = rl_canvas.Canvas(out, pagesize=(PW, PH))
    c.setTitle(slide_defs[0].get("title", "") if slide_defs else "")

    for i, d in enumerate(slide_defs):
        _RENDERERS[d["type"]](c, d)
        if i < len(slide_defs) - 1:
            c.showPage()

    c.save()
    print(f"✅ PDF生成完了: {out}  ({len(slide_defs)}ページ)")
    return out
