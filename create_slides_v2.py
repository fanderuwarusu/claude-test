"""
AI基礎講義_スライド_v2 生成スクリプト
- カラーテーマ A（ネイビー × ブルー）
- スライドサイズ 4:3（10.0" × 7.5"）
- パターン1: AI基礎講義（90分）全23枚
- v2: データ駆動型構造でスライド追加・変更が容易
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ============================================================
# カラーパレット（テーマA）
# ============================================================
C_PRIMARY = RGBColor(0x1A, 0x2E, 0x4A)
C_ACCENT  = RGBColor(0x00, 0x8C, 0xD7)
C_BG      = RGBColor(0xF0, 0xF6, 0xFC)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_GRAY    = RGBColor(0x6B, 0x7A, 0x8D)
C_PH      = RGBColor(0xCC, 0xCC, 0xCC)

W = Inches(10.0)
H = Inches(7.5)

FILENAME = "AI基礎講義_スライド_v2"


# ============================================================
# ユーティリティ
# ============================================================
def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, w, h,
             size=16, bold=False, color=C_PRIMARY,
             align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_ph_box(slide, x, y, w, h, label="※ここに内容を入力"):
    rect = slide.shapes.add_shape(1, x, y, w, h)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0xF7, 0xF7, 0xF7)
    rect.line.color.rgb = C_PH
    rect.line.width = Pt(1)
    tf = rect.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(12)
    run.font.color.rgb = C_PH
    run.font.italic = True
    return rect


def header(slide, title_text):
    add_rect(slide, Inches(0), Inches(0), W, Inches(1.0), C_PRIMARY)
    add_rect(slide, Inches(0), Inches(1.0), W, Inches(0.06), C_ACCENT)
    add_text(slide, title_text,
             Inches(0.35), Inches(0.15), Inches(9.3), Inches(0.75),
             size=22, bold=True, color=C_WHITE)


# ============================================================
# スライド描画関数
# ============================================================
def make_title_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_PRIMARY)
    add_rect(slide, Inches(0), Inches(0), Inches(0.45), H, C_ACCENT)
    add_text(slide, d["label"],
             Inches(0.7), Inches(1.4), Inches(8.8), Inches(0.55),
             size=14, bold=True, color=C_ACCENT)
    add_text(slide, d["title"],
             Inches(0.7), Inches(2.1), Inches(9.0), Inches(1.5),
             size=36, bold=True, color=C_WHITE)
    add_text(slide, d["subtitle"],
             Inches(0.7), Inches(3.75), Inches(9.0), Inches(0.8),
             size=18, color=RGBColor(0xA0, 0xC4, 0xE8))
    add_text(slide, "【日付】　【講師名】",
             Inches(0.7), Inches(5.6), Inches(7), Inches(0.45),
             size=13, color=C_GRAY)


def make_agenda_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_BG)
    header(slide, d["title"])
    y = Inches(1.18)
    row_h = Inches(0.72)
    for i, (num, label, dur) in enumerate(d["items"]):
        bg = C_ACCENT if i % 2 == 0 else C_PRIMARY
        add_rect(slide, Inches(0.35), y, Inches(0.6), row_h, bg)
        add_text(slide, num, Inches(0.35), y + Inches(0.15),
                 Inches(0.6), Inches(0.42),
                 size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_rect(slide, Inches(0.97), y, Inches(7.15), row_h,
                 RGBColor(0xFF, 0xFF, 0xFF))
        add_text(slide, label, Inches(1.07), y + Inches(0.15),
                 Inches(7.0), Inches(0.42), size=13, color=C_PRIMARY)
        add_rect(slide, Inches(8.15), y, Inches(1.5), row_h,
                 RGBColor(0xD8, 0xEE, 0xFA))
        add_text(slide, dur, Inches(8.15), y + Inches(0.15),
                 Inches(1.5), Inches(0.42),
                 size=12, color=C_ACCENT, align=PP_ALIGN.CENTER)
        y += row_h + Inches(0.04)


def make_section_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_ACCENT)
    add_rect(slide, Inches(0), Inches(0), W, Inches(0.07), C_WHITE)
    add_rect(slide, Inches(0), Inches(7.43), W, Inches(0.07), C_WHITE)
    add_text(slide, f"Section {d['num']}",
             Inches(0), Inches(2.3), W, Inches(0.65),
             size=18, color=C_WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, d["title"],
             Inches(0.5), Inches(2.95), Inches(9.0), Inches(1.1),
             size=32, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    if d.get("subtitle"):
        add_text(slide, d["subtitle"],
                 Inches(0.5), Inches(4.15), Inches(9.0), Inches(0.6),
                 size=16, color=RGBColor(0xD0, 0xEC, 0xFF),
                 align=PP_ALIGN.CENTER)


def make_content_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_BG)
    header(slide, d["title"])
    bullets = d.get("bullets")
    if bullets:
        tb = slide.shapes.add_textbox(Inches(0.4), Inches(1.2),
                                      Inches(9.2), Inches(5.5))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for b in bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(4)
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            if b.startswith("##"):
                run.text = b[2:].strip()
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = C_ACCENT
            else:
                run.text = f"▶  {b}"
                run.font.size = Pt(14)
                run.font.color.rgb = C_PRIMARY
    if d.get("ph"):
        ph_y = Inches(5.4) if bullets else Inches(1.3)
        ph_h = Inches(1.6) if bullets else Inches(5.5)
        add_ph_box(slide, Inches(0.4), ph_y, Inches(9.2), ph_h,
                   d.get("ph_label", "※ここに内容を入力"))
    if d.get("note"):
        add_text(slide, f"📌 {d['note']}",
                 Inches(0.4), Inches(7.05), Inches(9.2), Inches(0.35),
                 size=10, color=C_GRAY)


def make_two_col_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_BG)
    header(slide, d["title"])
    col_w = Inches(4.4)
    gap   = Inches(0.3)
    lx    = Inches(0.35)
    rx    = lx + col_w + gap
    add_rect(slide, lx, Inches(1.18), col_w, Inches(0.42), C_ACCENT)
    add_text(slide, d["left_title"],
             lx + Inches(0.08), Inches(1.2),
             col_w - Inches(0.16), Inches(0.38),
             size=13, bold=True, color=C_WHITE)
    add_rect(slide, rx, Inches(1.18), col_w, Inches(0.42), C_PRIMARY)
    add_text(slide, d["right_title"],
             rx + Inches(0.08), Inches(1.2),
             col_w - Inches(0.16), Inches(0.38),
             size=13, bold=True, color=C_WHITE)
    lb = d.get("left_bullets")
    rb = d.get("right_bullets")
    if lb:
        tb = slide.shapes.add_textbox(lx + Inches(0.1), Inches(1.72),
                                      col_w - Inches(0.2), Inches(5.3))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for b in lb:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(3)
            run = p.add_run()
            run.text = f"・{b}"
            run.font.size = Pt(13)
            run.font.color.rgb = C_PRIMARY
    else:
        add_ph_box(slide, lx, Inches(1.72), col_w, Inches(5.3))
    if rb:
        tb = slide.shapes.add_textbox(rx + Inches(0.1), Inches(1.72),
                                      col_w - Inches(0.2), Inches(5.3))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for b in rb:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(3)
            run = p.add_run()
            run.text = f"・{b}"
            run.font.size = Pt(13)
            run.font.color.rgb = C_PRIMARY
    else:
        add_ph_box(slide, rx, Inches(1.72), col_w, Inches(5.3))


def make_work_slide(prs, d):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, C_PRIMARY)
    add_rect(slide, Inches(0.4), Inches(0.3), Inches(1.7), Inches(0.55), C_ACCENT)
    add_text(slide, f"✏  {d.get('work_label', 'ワーク')}",
             Inches(0.45), Inches(0.32), Inches(1.6), Inches(0.5),
             size=13, bold=True, color=C_WHITE)
    add_text(slide, d["title"],
             Inches(0.4), Inches(1.1), Inches(9.2), Inches(1.0),
             size=26, bold=True, color=C_WHITE)
    add_text(slide, d["instruction"],
             Inches(0.4), Inches(2.2), Inches(9.2), Inches(0.65),
             size=15, color=RGBColor(0xA0, 0xC4, 0xE8))
    add_ph_box(slide, Inches(0.4), Inches(3.05), Inches(9.2), Inches(3.9),
               "※ワーク内容・手順をここに記載")


# ============================================================
# スライド定義（ここを編集してスライドを追加・変更）
# ============================================================
SLIDE_DEFS = [
    # ── 1: タイトル ───────────────────────────────────────────
    {
        "type": "title",
        "title": "AI基礎講義",
        "subtitle": "「AIを知る・使う」― 壁を下げ、業務で使える自分になる",
        "label": "Pattern 1  |  90分",
    },
    # ── 2: アジェンダ ─────────────────────────────────────────
    {
        "type": "agenda",
        "title": "本日のアジェンダ",
        "items": [
            ("0", "オープニング：今日のゴールと受講者のAI現在地確認",  "5分"),
            ("1", "生成AIとは何か：LLMの仕組みと限界",               "20分"),
            ("2", "AIツール実践：プロンプト設計＋業務ユースケース",   "40分"),
            ("3", "社会・産業へのAI浸透：競合・事例・動向",          "15分"),
            ("4", "リスクと注意点＋Q&A",                             "10分"),
        ],
    },
    # ── 3: 今日のゴール ───────────────────────────────────────
    {
        "type": "content",
        "title": "今日終わったら何が変わるか",
        "bullets": [
            "個人でAIを使っている状態　→　仕事でAIを使える状態へ",
            "AIを「怖い・難しい」から「使えるツール」に変える",
            "今日の終わりに：議事録・メール・要約を自分で試せる",
        ],
        "note": "オープニング  5分",
    },
    # ── 4: AI現在地確認 ───────────────────────────────────────
    {
        "type": "content",
        "title": "あなたの今のAI使用状況は？",
        "bullets": [
            "ChatGPT / Copilot / Gemini を使ったことがある？",
            "業務でAIを使っている？（週に何回くらい？）",
        ],
        "ph": True,
        "ph_label": "※ チェック項目 / アンケート結果スクリーンショットを貼り付け",
        "note": "オープニング  5分",
    },
    # ── 5: Section 1 ──────────────────────────────────────────
    {
        "type": "section",
        "num": 1,
        "title": "生成AIとは何か",
        "subtitle": "LLMの仕組みと限界を平易に理解する",
    },
    # ── 6: LLMの仕組み ────────────────────────────────────────
    {
        "type": "content",
        "title": "そもそもAIとは何か：LLMの仕組み",
        "bullets": [
            "LLM（大規模言語モデル）とは",
            "どうやって「文章を生成」しているか",
            "なぜ「賢く見える」のか",
        ],
        "ph": True,
        "ph_label": "※ 図解：LLMの仕組み（トークン予測・確率分布）を挿入",
        "note": "Section 1  |  20分",
    },
    # ── 7: ハルシネーション ───────────────────────────────────
    {
        "type": "content",
        "title": "なぜAIは間違えるのか：ハルシネーション",
        "bullets": [
            "「確率的な次の単語予測」という本質",
            "ハルシネーション（事実と異なる回答）が起きるメカニズム",
            "知識のカットオフ日とは",
            "→ だから「確認」が重要",
        ],
        "note": "Section 1  |  20分",
    },
    # ── 8: できること・できないこと ───────────────────────────
    {
        "type": "two_col",
        "title": "AIにできること・できないこと",
        "left_title": "✅ 得意なこと",
        "right_title": "❌ 苦手なこと",
        "left_bullets": [
            "文章の要約・生成・翻訳",
            "アイデア出し・ブレスト",
            "コード補完・デバッグ支援",
            "定型メール・報告書の下書き",
            "データ整形・変換のコード生成",
        ],
        "right_bullets": [
            "最新情報の取得（カットオフ以降）",
            "正確な数値計算・統計処理",
            "社内固有ルールの自律判断",
            "責任ある意思決定",
            "感情・文脈の深い読み取り",
        ],
    },
    # ── 9: ツール比較 ─────────────────────────────────────────
    {
        "type": "content",
        "title": "ChatGPT・Claude・Gemini、何が違うか",
        "bullets": [
            "## ChatGPT（OpenAI）",
            "最も普及。GPT-4o。Copilot との連携で企業利用拡大",
            "## Claude（Anthropic）",
            "長文処理・安全性重視。コンテキストウィンドウが広い",
            "## Gemini（Google）",
            "Google Workspace との統合。検索連動で最新情報に強い",
        ],
        "note": "Section 1  |  20分",
    },
    # ── 10: AIエコシステム全体像（v2追加） ────────────────────
    {
        "type": "content",
        "title": "AIエコシステムの全体像",
        "bullets": [
            "## 基盤モデル層",
            "OpenAI / Anthropic / Google / Meta（LLaMA）など",
            "## アプリケーション層",
            "業務SaaS（Notion AI・Salesforce Einstein・GitHub Copilot）",
            "## 社内活用層",
            "RAG（社内文書検索）・エージェント・RPAとの連携",
        ],
        "note": "Section 1  |  20分",
    },
    # ── 11: Section 2 ─────────────────────────────────────────
    {
        "type": "section",
        "num": 2,
        "title": "AIツール実践",
        "subtitle": "プロンプト設計＋業務ユースケース（40分）",
    },
    # ── 12: 良い・悪いプロンプト ──────────────────────────────
    {
        "type": "two_col",
        "title": "良いプロンプトと悪いプロンプト",
        "left_title": "❌ 悪い例",
        "right_title": "✅ 良い例",
        "left_bullets": [
            "「まとめて」",
            "「なんかいい感じに」",
            "「報告書を書いて」",
            "",
            "→ 曖昧・ゴールが不明",
        ],
        "right_bullets": [
            "役割を指定する",
            "背景・条件を明示する",
            "出力形式を指定する",
            "例を示す（Few-shot）",
            "→ 具体・制約・ゴールが明確",
        ],
    },
    # ── 13: プロンプトテンプレート集（v2追加） ────────────────
    {
        "type": "content",
        "title": "すぐ使えるプロンプトテンプレート",
        "bullets": [
            "## 要約テンプレート",
            "「以下の文章を[対象読者]向けに[字数]字以内で要約してください」",
            "## 文章生成テンプレート",
            "「あなたは[役割]です。[背景]を踏まえて[形式]で書いてください」",
            "## 壁打ちテンプレート",
            "「[アイデア]について、メリット・デメリット・改善案を挙げてください」",
        ],
        "note": "Section 2  |  40分",
    },
    # ── 14: ライブデモ①議事録 ────────────────────────────────
    {
        "type": "content",
        "title": "ライブデモ①：議事録の自動要約",
        "bullets": [
            "手順：音声テキスト or 議事録テキストをそのまま貼り付ける",
            "プロンプト例：「以下の議事録を要約してください。決定事項・TODO・懸念点を箇条書きで」",
        ],
        "ph": True,
        "ph_label": "※ ライブデモ画面キャプチャ / プロンプトテンプレートを挿入",
        "note": "Section 2  |  40分",
    },
    # ── 15: ライブデモ②メール ────────────────────────────────
    {
        "type": "content",
        "title": "ライブデモ②：メール文案の生成",
        "bullets": [
            "手順：背景・相手・目的をひとこと伝えるだけ",
            "プロンプト例：「社内の営業部に依頼するメールを丁寧なビジネス文体で書いて」",
        ],
        "ph": True,
        "ph_label": "※ ライブデモ画面キャプチャ / プロンプトテンプレートを挿入",
        "note": "Section 2  |  40分",
    },
    # ── 16: ライブデモ③データ要約（v2追加） ──────────────────
    {
        "type": "content",
        "title": "ライブデモ③：表データの要約・考察",
        "bullets": [
            "手順：CSVや表をテキスト形式で貼り付ける",
            "プロンプト例：「以下のデータを分析し、傾向・課題・改善提案を3点ずつ挙げてください」",
        ],
        "ph": True,
        "ph_label": "※ ライブデモ画面キャプチャ / プロンプトテンプレートを挿入",
        "note": "Section 2  |  40分",
    },
    # ── 17: ワーク ────────────────────────────────────────────
    {
        "type": "work",
        "title": "自分でプロンプトを1つ打ってみよう",
        "instruction": "今日学んだ「良いプロンプトの型」を使って、実際に1問送信しましょう（5分）",
        "work_label": "ワーク（5分）",
    },
    # ── 18: Section 3 ─────────────────────────────────────────
    {
        "type": "section",
        "num": 3,
        "title": "社会・産業へのAI浸透",
        "subtitle": "業種別事例と「他社はここまでやっている」",
    },
    # ── 19: 業種別事例 ────────────────────────────────────────
    {
        "type": "content",
        "title": "業種別 AI導入事例",
        "bullets": [
            "## 事例①　【受講者企業に近い業種】",
            "※事例内容・効果をここに記載",
            "## 事例②　【受講者企業に近い業種】",
            "※事例内容・効果をここに記載",
            "## 事例③　【受講者企業に近い業種】",
            "※事例内容・効果をここに記載",
        ],
        "note": "Section 3  |  15分  ｜  受講者企業の業種に合わせて差し替えること",
    },
    # ── 20: AI導入の効果・ROI（v2追加） ──────────────────────
    {
        "type": "content",
        "title": "AI導入で何が変わるか：生産性・ROIの実態",
        "bullets": [
            "## 時間削減効果（McKinsey 2024調査より）",
            "定型業務（メール・要約・調査）：週平均 約3〜5時間削減",
            "## コスト試算例",
            "10人チーム × 年間240時間削減 ≒ 年間約480万円相当の効率化",
            "## 導入成功の共通点",
            "小さく始めて効果を測る → 横展開する",
        ],
        "note": "Section 3  |  15分",
    },
    # ── 21: Section 4 ─────────────────────────────────────────
    {
        "type": "section",
        "num": 4,
        "title": "リスクと注意点",
        "subtitle": "安全に使うための最低限の知識",
    },
    # ── 22: リスクと注意点 ────────────────────────────────────
    {
        "type": "content",
        "title": "AIを使う上でのリスクと注意点",
        "bullets": [
            "情報漏洩リスク：個人情報・社外秘は入力しない",
            "ハルシネーション：出力は必ず確認・ファクトチェック",
            "著作権・利用規約：商用利用時は各サービスの規約を確認",
            "社内ルールの確認：利用可能なツール・用途を把握しておく",
        ],
        "note": "Section 4  |  10分",
    },
    # ── 23: まとめ・Q&A ───────────────────────────────────────
    {
        "type": "content",
        "title": "まとめ・Q&A",
        "bullets": [
            "## 今日の3つのポイント",
            "① AIは「確率的な文章予測」――必ず確認する習慣を",
            "② 良いプロンプトは「役割・条件・形式」を明示する",
            "③ 他社はすでに動いている――まず1つ業務で試してみる",
        ],
        "ph": True,
        "ph_label": "※ Q&Aメモ欄",
        "note": "10分",
    },
]

# ============================================================
# ディスパッチ & ビルド
# ============================================================
RENDERERS = {
    "title":   make_title_slide,
    "agenda":  make_agenda_slide,
    "section": make_section_slide,
    "content": make_content_slide,
    "two_col": make_two_col_slide,
    "work":    make_work_slide,
}


def build():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    for d in SLIDE_DEFS:
        RENDERERS[d["type"]](prs, d)

    out = f"{FILENAME}.pptx"
    prs.save(out)
    print(f"✅ PPTX生成完了: {out}  ({len(prs.slides)}枚)")
    return out


if __name__ == "__main__":
    build()
