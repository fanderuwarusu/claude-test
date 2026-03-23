"""
PPTX → PDF 変換スクリプト（python-pptx + reportlab使用）
スライドの内容をテキストとして抽出しPDFに出力する
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit
import os, glob

# 日本語フォント登録
FONT_PATHS = [
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
    "/usr/share/fonts/truetype/takao-gothic/TakaoGothic.ttf",
    "/usr/share/fonts/truetype/vlgothic/VL-Gothic-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]

JP_FONT = "Helvetica"  # fallback
for fp in FONT_PATHS:
    if os.path.exists(fp):
        try:
            pdfmetrics.registerFont(TTFont("JPFont", fp))
            JP_FONT = "JPFont"
            print(f"Using font: {fp}")
            break
        except Exception as e:
            continue

PAGE_W, PAGE_H = landscape(A4)

COLOR_NAVY   = colors.HexColor("#1A2E4A")
COLOR_ACCENT = colors.HexColor("#008CD7")
COLOR_LIGHT  = colors.HexColor("#F0F6FC")
COLOR_WHITE  = colors.white
COLOR_GRAY   = colors.HexColor("#6B7A8D")
COLOR_PH     = colors.HexColor("#CCCCCC")


def get_slide_bg_color(slide):
    """スライド背景色の判定（タイトル系か通常か）"""
    try:
        fill = slide.background.fill
        if fill.type is not None:
            rgb = fill.fore_color.rgb
            if rgb == RGBColor(0x1A, 0x2E, 0x4A):
                return "navy"
            elif rgb == RGBColor(0x00, 0x8C, 0xD7):
                return "accent"
        return "light"
    except:
        return "light"


def extract_texts(slide):
    """スライドからテキストを抽出"""
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                line = para.text.strip()
                if line:
                    font_size = None
                    bold = False
                    color_rgb = None
                    for run in para.runs:
                        if run.font.size:
                            font_size = run.font.size.pt
                        if run.font.bold:
                            bold = True
                        try:
                            color_rgb = run.font.color.rgb
                        except:
                            pass
                    texts.append({
                        "text": line,
                        "size": font_size or 12,
                        "bold": bold,
                        "color": color_rgb,
                    })
    return texts


def draw_slide(c, slide, page_num):
    bg = get_slide_bg_color(slide)

    # 背景
    if bg == "navy":
        c.setFillColor(COLOR_NAVY)
    elif bg == "accent":
        c.setFillColor(COLOR_ACCENT)
    else:
        c.setFillColor(COLOR_LIGHT)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)

    # アクセントバー（左または上）
    if bg in ("navy", "accent"):
        c.setFillColor(COLOR_ACCENT if bg == "navy" else COLOR_WHITE)
        c.rect(0, 0, 14, PAGE_H, fill=1, stroke=0)
    else:
        c.setFillColor(COLOR_NAVY)
        c.rect(0, PAGE_H - 55, PAGE_W, 55, fill=1, stroke=0)
        c.setFillColor(COLOR_ACCENT)
        c.rect(0, PAGE_H - 58, PAGE_W, 4, fill=1, stroke=0)

    texts = extract_texts(slide)

    y = PAGE_H - 40 if bg in ("navy", "accent") else PAGE_H - 50
    first = True

    for t in texts:
        text = t["text"]
        size = min(t["size"], 36)
        bold = t["bold"]

        # 色決定
        if bg in ("navy", "accent"):
            txt_color = COLOR_WHITE
            if t["color"]:
                try:
                    r, g, b = t["color"]
                    if (r, g, b) == (0, 140, 215):  # accent
                        txt_color = colors.HexColor("#A0C4E8")
                    elif (r, g, b) == (107, 122, 141):
                        txt_color = colors.HexColor("#8898AA")
                except:
                    pass
        else:
            if first and size >= 20:
                txt_color = COLOR_WHITE
            else:
                txt_color = COLOR_NAVY
                if t["color"]:
                    try:
                        r, g, b = t["color"]
                        if (r, g, b) == (0, 140, 215):
                            txt_color = COLOR_ACCENT
                        elif (r, g, b) == (204, 204, 204):
                            txt_color = COLOR_PH
                    except:
                        pass

        # プレースホルダーテキストはグレーに
        if "※" in text or "【" in text:
            txt_color = COLOR_PH if bg == "light" else colors.HexColor("#8898AA")
            size = min(size, 12)

        c.setFillColor(txt_color)
        font_size = max(8, min(size, 28))
        c.setFont(JP_FONT, font_size)

        x_start = 30 if bg in ("navy", "accent") else 25
        max_w = PAGE_W - x_start - 20

        lines = simpleSplit(text, JP_FONT, font_size, max_w)
        for line in lines:
            if y < 20:
                break
            c.drawString(x_start, y, line)
            y -= font_size * 1.3

        first = False
        y -= 4

    # ページ番号
    c.setFillColor(COLOR_GRAY)
    c.setFont(JP_FONT, 8)
    c.drawRightString(PAGE_W - 10, 8, str(page_num))


def convert(pptx_path, pdf_path):
    prs = Presentation(pptx_path)
    c = canvas.Canvas(pdf_path, pagesize=landscape(A4))
    for i, slide in enumerate(prs.slides, 1):
        draw_slide(c, slide, i)
        c.showPage()
    c.save()
    print(f"✅ {pdf_path} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    convert("AI基礎講義_スライド.pptx",             "AI基礎講義_スライド.pdf")
    convert("AI業務導入ワークショップ_スライド.pptx", "AI業務導入ワークショップ_スライド.pdf")
